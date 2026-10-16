from __future__ import annotations

import json
import re
import shutil
import zipfile
from pathlib import Path
from typing import Any, Iterable

from docx import Document

from .models import FieldResolution
from .utils import normalize_text


PENDING = "PENDING VERIFICATION"


class TemplateEngine:
    """
    Preserves the original DOCX/ODT package and replaces only configured placeholders
    or configured sample values. For best reliability, put placeholders such as
    {{borrower_name}} directly in the master template.
    """

    def build_values(
        self,
        resolutions: dict[str, FieldResolution],
        controlled_draft: bool = True,
    ) -> dict[str, str]:
        values: dict[str, str] = {}
        for field_name, resolution in resolutions.items():
            if resolution.status in {"conflict", "review_required", "missing"} and controlled_draft:
                values[field_name] = PENDING
            elif resolution.selected_value is None:
                values[field_name] = PENDING if controlled_draft else ""
            else:
                values[field_name] = self._format_value(resolution.selected_value)
        return values

    def generate(
        self,
        template_path: Path,
        output_path: Path,
        values: dict[str, str],
        template_map_path: Path | None = None,
        leakage_terms: list[str] | None = None,
    ) -> dict[str, Any]:
        output_path.parent.mkdir(parents=True, exist_ok=True)
        suffix = template_path.suffix.lower()
        template_map = self._load_template_map(template_map_path)

        if suffix == ".docx":
            self._generate_docx(template_path, output_path, values, template_map)
            rendered_text = self._extract_docx_text(output_path)
        elif suffix == ".odt":
            self._generate_odt(template_path, output_path, values, template_map)
            rendered_text = self._extract_odt_text(output_path)
        else:
            raise ValueError("Only DOCX and ODT templates are supported for structure-preserving generation.")

        remaining_placeholders = sorted(set(re.findall(r"\{\{\s*([A-Za-z0-9_]+)\s*\}\}", rendered_text)))
        leakage_hits = []
        for term in leakage_terms or []:
            if term and normalize_text(term) in normalize_text(rendered_text):
                leakage_hits.append(term)

        return {
            "output": str(output_path),
            "remaining_placeholders": remaining_placeholders,
            "leakage_hits": leakage_hits,
            "passed": not remaining_placeholders and not leakage_hits,
        }

    @staticmethod
    def _format_value(value: Any) -> str:
        if isinstance(value, float):
            return f"{value:g}"
        if isinstance(value, (dict, list)):
            return json.dumps(value, ensure_ascii=False)
        return str(value)

    @staticmethod
    def _load_template_map(path: Path | None) -> dict[str, str]:
        if path is None:
            return {}
        with path.open("r", encoding="utf-8") as fh:
            data = json.load(fh)
        return {str(k): str(v) for k, v in data.get("exact_text_to_field", {}).items()}

    def _replacement_pairs(self, values: dict[str, str], template_map: dict[str, str]) -> list[tuple[str, str]]:
        pairs = [(f"{{{{{field_name}}}}}", value) for field_name, value in values.items()]
        for old_text, field_name in template_map.items():
            pairs.append((old_text, values.get(field_name, PENDING)))
        return sorted(pairs, key=lambda pair: len(pair[0]), reverse=True)

    def _generate_docx(
        self,
        template: Path,
        output: Path,
        values: dict[str, str],
        template_map: dict[str, str],
    ) -> None:
        doc = Document(template)
        pairs = self._replacement_pairs(values, template_map)
        for paragraph in self._all_docx_paragraphs(doc):
            self._replace_in_paragraph(paragraph, pairs)
        doc.save(output)

    def _all_docx_paragraphs(self, doc: Document) -> Iterable:
        for paragraph in doc.paragraphs:
            yield paragraph
        for table in doc.tables:
            yield from self._table_paragraphs(table)
        for section in doc.sections:
            for paragraph in section.header.paragraphs:
                yield paragraph
            for table in section.header.tables:
                yield from self._table_paragraphs(table)
            for paragraph in section.footer.paragraphs:
                yield paragraph
            for table in section.footer.tables:
                yield from self._table_paragraphs(table)

    def _table_paragraphs(self, table) -> Iterable:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    yield paragraph
                for nested in cell.tables:
                    yield from self._table_paragraphs(nested)

    @staticmethod
    def _replace_in_paragraph(paragraph, pairs: list[tuple[str, str]]) -> None:
        if not paragraph.runs:
            return
        combined = "".join(run.text for run in paragraph.runs)
        replaced = combined
        for old, new in pairs:
            replaced = replaced.replace(old, new)
        if replaced == combined:
            return
        paragraph.runs[0].text = replaced
        for run in paragraph.runs[1:]:
            run.text = ""

    def _generate_odt(
        self,
        template: Path,
        output: Path,
        values: dict[str, str],
        template_map: dict[str, str],
    ) -> None:
        shutil.copy2(template, output)
        pairs = self._replacement_pairs(values, template_map)
        temp_output = output.with_suffix(output.suffix + ".tmp")
        with zipfile.ZipFile(output, "r") as source, zipfile.ZipFile(temp_output, "w") as target:
            for item in source.infolist():
                data = source.read(item.filename)
                if item.filename in {"content.xml", "styles.xml"}:
                    text = data.decode("utf-8")
                    for old, new in pairs:
                        text = text.replace(self._xml_escape(old), self._xml_escape(new))
                        text = text.replace(old, self._xml_escape(new))
                    data = text.encode("utf-8")
                target.writestr(item, data)
        temp_output.replace(output)

    @staticmethod
    def _xml_escape(value: str) -> str:
        return value.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    def _extract_docx_text(self, path: Path) -> str:
        doc = Document(path)
        return "\n".join(p.text for p in self._all_docx_paragraphs(doc))

    @staticmethod
    def _extract_odt_text(path: Path) -> str:
        with zipfile.ZipFile(path) as archive:
            data = archive.read("content.xml").decode("utf-8", errors="replace")
        return re.sub(r"<[^>]+>", " ", data)
