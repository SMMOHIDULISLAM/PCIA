from __future__ import annotations

import json
import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document

from pcia.pipeline import ProposalPipeline
from pcia.validate import validate_candidate


class PipelineTests(unittest.TestCase):
    def test_rejects_form_labels_as_names(self):
        for value in ["No. and", "Person Contacted", "Company Structure", "Year in Present Business"]:
            valid, reason, _, _ = validate_candidate("borrower_name", value)
            self.assertFalse(valid, value)
            self.assertTrue(reason)

    def test_processes_zip_and_resolves_core_fields(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            source = root / "docs"
            source.mkdir()
            (source / "Forwarding.txt").write_text(
                """
                CREDIT PROPOSAL
                Branch: Jagannathpur Branch, Sunamgonj
                Name of Applicant: Rana Miah
                Nature of Facility: Pubali Abashon Loan
                Loan Amount: Tk.50,00,000
                Purpose: Completion of civil and finishing works of residential building
                Tenure: 5 years
                Grace Period: 6 months
                Rate of Interest: 12%
                """,
                encoding="utf-8",
            )
            proposal_zip = root / "proposal.zip"
            with zipfile.ZipFile(proposal_zip, "w", zipfile.ZIP_DEFLATED) as z:
                z.write(source / "Forwarding.txt", "Forwarding.txt")

            pipeline = ProposalPipeline(root / "cases", enable_ocr=False)
            result = pipeline.process(proposal_zip)
            self.assertEqual(result.resolutions["borrower_name"].selected_value, "Rana Miah")
            self.assertIn("Jagannathpur", result.resolutions["branch"].selected_value)
            self.assertEqual(result.resolutions["loan_product"].selected_value, "Pubali Abashon Loan")
            self.assertEqual(result.resolutions["loan_amount"].selected_value, 5_000_000.0)
            self.assertEqual(result.resolutions["tenure_months"].selected_value, 60)
            self.assertTrue((result.case_dir / f"{result.case_id}_results.zip").exists())

    def test_docx_placeholder_replacement_and_leakage_gate(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            template = root / "template.docx"
            output = root / "output.docx"
            doc = Document()
            doc.add_paragraph("Borrower: {{borrower_name}}")
            doc.add_paragraph("Old borrower: Ali Ahmed")
            doc.save(template)

            from pcia.template_engine import TemplateEngine
            engine = TemplateEngine()
            quality = engine.generate(
                template,
                output,
                {"borrower_name": "Rana Miah"},
                leakage_terms=["Ali Ahmed"],
            )
            self.assertIn("Ali Ahmed", quality["leakage_hits"])
            self.assertFalse(quality["passed"])


if __name__ == "__main__":
    unittest.main()
