from __future__ import annotations

import csv
import json
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path
from typing import Iterable, Any
from xml.etree import ElementTree as ET

import fitz
from docx import Document
from openpyxl import load_workbook

from .models import DocumentPage, ManifestItem
from .utils import normalize_document_text, normalize_spaces


ODF_NS = {
    "office": "urn:oasis:names:tc:opendocument:xmlns:office:1.0",
    "text": "urn:oasis:names:tc:opendocument:xmlns:text:1.0",
    "table": "urn:oasis:names:tc:opendocument:xmlns:table:1.0",
}


class DocumentReader:
    def __init__(self, enable_ocr: bool = False, ocr_language: str = "eng") -> None:
        self.enable_ocr = enable_ocr
        self.ocr_language = ocr_language

    def read_all(self, source_dir: Path, manifest: list[ManifestItem]) -> list[DocumentPage]:
        pages: list[DocumentPage] = []
        manifest_index = {x.relative_path: x for x in manifest}
        for file_path in sorted(p for p in source_dir.rglob("*") if p.is_file()):
            relative = file_path.relative_to(source_dir).as_posix()
            item = manifest_index[relative]
            try:
                extracted = list(self._read_file(file_path, relative))
                pages.extend(extracted)
                item.page_count = len(extracted)
                item.status = "processed" if extracted else "unsupported_or_empty"
                if not extracted:
                    item.message = "No text-bearing pages produced."
            except Exception as exc:
                item.status = "failed"
                item.message = f"{type(exc).__name__}: {exc}"
        return pages

    def _read_file(self, path: Path, relative: str) -> Iterable[DocumentPage]:
        suffix = path.suffix.lower()
        if suffix == ".pdf":
            yield from self._read_pdf(path, relative)
        elif suffix == ".docx":
            yield self._read_docx(path, relative)
        elif suffix == ".odt":
            yield self._read_odt(path, relative)
        elif suffix in {".xlsx", ".xlsm"}:
            yield from self._read_xlsx(path, relative)
        elif suffix == ".ods":
            yield from self._read_ods(path, relative)
        elif suffix in {".txt", ".csv", ".json"}:
            yield self._read_textual(path, relative)
        elif suffix in {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp"}:
            yield self._read_image(path, relative)
        elif suffix in {".doc", ".xls"}:
            converted = self._convert_legacy(path)
            if converted is not None:
                yield from self._read_file(converted, relative + f" [converted:{converted.suffix}]")

    def _read_pdf(self, path: Path, relative: str) -> Iterable[DocumentPage]:
        doc = fitz.open(path)
        try:
            for index in range(doc.page_count):
                page = doc.load_page(index)
                native_text = page.get_text("text") or ""
                text = normalize_document_text(native_text)
                method = "native_pdf_text"
                if self.enable_ocr and len(text) < 40:
                    ocr_text = self._ocr_pdf_page(page)
                    if len(ocr_text) > len(text):
                        text = ocr_text
                        method = "ocr"

                tables = self._extract_pdf_tables(page)
                table_lines = self._tables_to_lines(tables)
                if table_lines:
                    text = normalize_document_text(text + "\n" + "\n".join(table_lines))

                yield DocumentPage(
                    relative_path=relative,
                    page_number=index + 1,
                    text=text,
                    document_type="unclassified",
                    extraction_method=method + ("+tables" if tables else ""),
                    readable=bool(text.strip()),
                    metadata={
                        "width": page.rect.width,
                        "height": page.rect.height,
                        "tables": tables,
                    },
                )
        finally:
            doc.close()

    @staticmethod
    def _extract_pdf_tables(page: fitz.Page) -> list[list[list[str]]]:
        tables: list[list[list[str]]] = []
        try:
            finder = page.find_tables()
            for table in finder.tables:
                rows: list[list[str]] = []
                for row in table.extract():
                    cleaned = [normalize_spaces(str(cell or "")) for cell in row]
                    if any(cleaned):
                        rows.append(cleaned)
                if rows:
                    tables.append(rows)
        except Exception:
            # Table extraction is supplementary; native text must still survive.
            pass
        return tables

    @staticmethod
    def _tables_to_lines(tables: list[list[list[str]]]) -> list[str]:
        lines: list[str] = []
        for table_index, table in enumerate(tables, start=1):
            lines.append(f"TABLE {table_index}")
            for row in table:
                lines.append("TABLE_ROW | " + " | ".join(row))
        return lines

    def _ocr_pdf_page(self, page: fitz.Page) -> str:
        try:
            import pytesseract
            from PIL import Image
        except ImportError:
            return ""
        pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
        image = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
        return normalize_document_text(pytesseract.image_to_string(image, lang=self.ocr_language))

    def _read_docx(self, path: Path, relative: str) -> DocumentPage:
        doc = Document(path)
        blocks: list[str] = []
        for section in doc.sections:
            blocks.extend(p.text for p in section.header.paragraphs if p.text.strip())
            for table in section.header.tables:
                blocks.extend(self._table_text(table))
        blocks.extend(p.text for p in doc.paragraphs if p.text.strip())
        for table in doc.tables:
            blocks.extend(self._table_text(table))
        for section in doc.sections:
            blocks.extend(p.text for p in section.footer.paragraphs if p.text.strip())
            for table in section.footer.tables:
                blocks.extend(self._table_text(table))
        return DocumentPage(relative, 1, normalize_document_text("\n".join(blocks)), "unclassified", "docx", bool(blocks))

    @staticmethod
    def _table_text(table) -> list[str]:
        rows: list[str] = []
        for row in table.rows:
            cells = [normalize_spaces(cell.text) for cell in row.cells]
            rows.append("TABLE_ROW | " + " | ".join(cells))
        return rows

    def _read_odt(self, path: Path, relative: str) -> DocumentPage:
        with zipfile.ZipFile(path) as archive:
            root = ET.fromstring(archive.read("content.xml"))
        blocks: list[str] = []
        tables: list[list[list[str]]] = []

        for elem in root.iter():
            local = elem.tag.rsplit("}", 1)[-1]
            if local in {"p", "h"}:
                text = normalize_spaces("".join(elem.itertext()))
                if text:
                    blocks.append(text)

        for table_elem in root.findall(".//table:table", ODF_NS):
            table_rows = self._odf_table_rows(table_elem)
            if table_rows:
                tables.append(table_rows)
        blocks.extend(self._tables_to_lines(tables))
        text = normalize_document_text("\n".join(blocks))
        return DocumentPage(relative, 1, text, "unclassified", "odt_xml", bool(text), {"tables": tables})

    def _read_ods(self, path: Path, relative: str) -> Iterable[DocumentPage]:
        with zipfile.ZipFile(path) as archive:
            root = ET.fromstring(archive.read("content.xml"))
        sheets = root.findall(".//table:table", ODF_NS)
        for index, sheet in enumerate(sheets, start=1):
            sheet_name = sheet.attrib.get(f"{{{ODF_NS['table']}}}name", f"Sheet{index}")
            rows = self._odf_table_rows(sheet)
            lines = [f"SHEET: {sheet_name}"] + ["TABLE_ROW | " + " | ".join(row) for row in rows]
            text = normalize_document_text("\n".join(lines))
            yield DocumentPage(
                relative,
                index,
                text,
                "unclassified",
                "ods_xml",
                bool(text),
                {"sheet": sheet_name, "tables": [rows]},
            )

    @staticmethod
    def _odf_table_rows(table_elem: ET.Element) -> list[list[str]]:
        rows: list[list[str]] = []
        row_repeat_attr = f"{{{ODF_NS['table']}}}number-rows-repeated"
        col_repeat_attr = f"{{{ODF_NS['table']}}}number-columns-repeated"
        value_attr = f"{{{ODF_NS['office']}}}value"
        string_value_attr = f"{{{ODF_NS['office']}}}string-value"

        for row_elem in table_elem.findall("./table:table-row", ODF_NS):
            row: list[str] = []
            for cell in list(row_elem):
                if cell.tag.rsplit("}", 1)[-1] not in {"table-cell", "covered-table-cell"}:
                    continue
                text_parts = []
                for p in cell.findall(".//text:p", ODF_NS):
                    val = normalize_spaces("".join(p.itertext()))
                    if val:
                        text_parts.append(val)
                cell_text = " ".join(text_parts)
                if not cell_text:
                    cell_text = cell.attrib.get(string_value_attr, cell.attrib.get(value_attr, ""))
                repeat = min(int(cell.attrib.get(col_repeat_attr, "1") or "1"), 100)
                row.extend([normalize_spaces(str(cell_text))] * repeat)
            while row and not row[-1]:
                row.pop()
            repeat_rows = min(int(row_elem.attrib.get(row_repeat_attr, "1") or "1"), 100)
            if any(row):
                rows.extend([row.copy() for _ in range(repeat_rows)])
        return rows

    def _read_xlsx(self, path: Path, relative: str) -> Iterable[DocumentPage]:
        workbook = load_workbook(path, data_only=False, read_only=True)
        for index, sheet in enumerate(workbook.worksheets, start=1):
            rows: list[list[str]] = []
            for row in sheet.iter_rows(values_only=True):
                values = ["" if value is None else str(value) for value in row]
                if any(v.strip() for v in values):
                    rows.append(values)
            lines = [f"SHEET: {sheet.title}"] + ["TABLE_ROW | " + " | ".join(row) for row in rows]
            text = normalize_document_text("\n".join(lines))
            yield DocumentPage(relative, index, text, "unclassified", "xlsx", bool(text), {"sheet": sheet.title, "tables": [rows]})

    def _read_textual(self, path: Path, relative: str) -> DocumentPage:
        if path.suffix.lower() == ".csv":
            with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as fh:
                rows = ["TABLE_ROW | " + " | ".join(row) for row in csv.reader(fh)]
            text = "\n".join(rows)
        elif path.suffix.lower() == ".json":
            data = json.loads(path.read_text(encoding="utf-8-sig", errors="replace"))
            text = json.dumps(data, ensure_ascii=False, indent=2)
        else:
            text = path.read_text(encoding="utf-8-sig", errors="replace")
        text = normalize_document_text(text)
        return DocumentPage(relative, 1, text, "unclassified", "text", bool(text))

    def _read_image(self, path: Path, relative: str) -> DocumentPage:
        text = ""
        if self.enable_ocr:
            try:
                import pytesseract
                from PIL import Image
                text = pytesseract.image_to_string(Image.open(path), lang=self.ocr_language)
            except Exception:
                text = ""
        text = normalize_document_text(text)
        return DocumentPage(relative, 1, text, "unclassified", "ocr" if text else "image", bool(text))

    @staticmethod
    def _convert_legacy(path: Path) -> Path | None:
        office = shutil.which("libreoffice") or shutil.which("soffice")
        if not office:
            return None
        output_dir = Path(tempfile.mkdtemp(prefix="pcia-convert-"))
        target_format = "docx" if path.suffix.lower() == ".doc" else "xlsx"
        command = [office, "--headless", "--convert-to", target_format, "--outdir", str(output_dir), str(path)]
        process = subprocess.run(command, capture_output=True, text=True, timeout=120)
        if process.returncode != 0:
            return None
        converted = output_dir / (path.stem + "." + target_format)
        return converted if converted.exists() else None
