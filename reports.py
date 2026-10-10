from __future__ import annotations

import csv
import json
import zipfile
from dataclasses import asdict
from pathlib import Path
from typing import Any

from docx import Document
from docx.shared import Pt

from .models import ProcessingResult
from .utils import dump_json


class ReportWriter:
    def write_all(self, result: ProcessingResult) -> list[Path]:
        output_dir = result.case_dir / "outputs"
        output_dir.mkdir(parents=True, exist_ok=True)

        files: list[Path] = []
        files.append(self._json(output_dir / "package_manifest.json", [asdict(x) for x in result.manifest]))
        files.append(self._json(output_dir / "evidence_ledger.json", [asdict(x) for x in result.candidates]))
        files.append(self._json(output_dir / "resolved_case.json", {k: asdict(v) for k, v in result.resolutions.items()}))
        files.append(self._json(output_dir / "conflicts_and_anomalies.json", result.anomalies))
        files.append(self._json(output_dir / "operator_questions.json", result.operator_questions))
        files.append(self._json(output_dir / "variation_explanations.json", result.variation_explanations))
        files.append(self._json(output_dir / "missing_fields.json", result.missing_fields))
        files.append(self._json(output_dir / "missing_documents.json", result.missing_documents))
        files.append(self._write_evidence_csv(output_dir / "evidence_ledger.csv", result))
        files.append(self._write_summary(output_dir / "SUMMARY.md", result))
        files.append(self._write_anomaly_letter(output_dir / "Anomaly_and_Shortcoming_Letter_Draft.docx", result))
        return files

    @staticmethod
    def _json(path: Path, data: Any) -> Path:
        dump_json(path, data)
        return path

    @staticmethod
    def _write_evidence_csv(path: Path, result: ProcessingResult) -> Path:
        with path.open("w", encoding="utf-8-sig", newline="") as fh:
            writer = csv.writer(fh)
            writer.writerow([
                "field_name", "value", "valid", "rejection_reason", "document_type",
                "file", "page", "confidence", "source_priority", "pattern_id",
            ])
            for item in result.candidates:
                writer.writerow([
                    item.field_name,
                    json.dumps(item.value, ensure_ascii=False) if isinstance(item.value, (dict, list)) else item.value,
                    item.valid,
                    item.rejection_reason,
                    item.document_type,
                    item.relative_path,
                    item.page_number,
                    item.confidence,
                    item.source_priority,
                    item.pattern_id,
                ])
        return path

    @staticmethod
    def _write_summary(path: Path, result: ProcessingResult) -> Path:
        def value(field: str) -> str:
            resolution = result.resolutions.get(field)
            return str(resolution.selected_value) if resolution and resolution.selected_value is not None else "PENDING VERIFICATION"

        lines = [
            f"# PCIA Case Summary: {value('borrower_name')}",
            "",
            f"- Case ID: `{result.case_id}`",
            f"- Borrower type: **{result.borrower_type}**",
            f"- Branch: **{value('branch')}**",
            f"- Product: **{value('loan_product')}**",
            f"- Amount: **{value('loan_amount')}**",
            f"- Purpose: **{value('purpose')}**",
            f"- Files inventoried: **{len(result.manifest)}**",
            f"- Pages/sections read: **{len(result.pages)}**",
            f"- Missing fields: **{len(result.missing_fields)}**",
            f"- Missing document types: **{len(result.missing_documents)}**",
            f"- Anomalies/conflicts: **{len(result.anomalies)}**",
            f"- Operator questions: **{len(result.operator_questions)}**",
            f"- Explained acceptable variations: **{len(result.variation_explanations)}**",
            "",
            "## Generation rule",
            "Final documents must not be issued while critical conflicts remain. Controlled drafts may use `PENDING VERIFICATION`; old template-borrower data must never remain.",
            "",
            "## Variation handling",
            "Account formatting, Mia/Miah spellings, old/smart NID lengths, and general-versus-recommended rates are explained separately and are not counted as anomalies unless the underlying fact truly differs.",
        ]
        path.write_text("\n".join(lines), encoding="utf-8")
        return path

    @staticmethod
    def _write_anomaly_letter(path: Path, result: ProcessingResult) -> Path:
        def value(field: str) -> str:
            resolution = result.resolutions.get(field)
            return str(resolution.selected_value) if resolution and resolution.selected_value is not None else "PENDING VERIFICATION"

        doc = Document()
        normal = doc.styles["Normal"]
        normal.font.name = "Times New Roman"
        normal.font.size = Pt(11)
        p = doc.add_paragraph()
        p.alignment = 1
        run = p.add_run("PUBALI BANK PLC\nREGIONAL OFFICE, SYLHET WEST")
        run.bold = True
        doc.add_paragraph("Date: SYSTEM DATE")
        doc.add_paragraph("Reference: SYSTEM-GENERATED BANK REFERENCE")
        doc.add_paragraph("The Branch Manager\n" + value("branch"))
        doc.add_paragraph(
            f"Subject: Anomalies and shortcomings in the credit proposal of {value('borrower_name')} - {value('loan_product')}"
        ).runs[0].bold = True
        doc.add_paragraph("Dear Sir,")
        doc.add_paragraph(
            f"With reference to your credit proposal of {value('borrower_name')} under {value('loan_product')}, "
            "we found the following anomalies and shortcomings. Please rectify and submit the supporting documents."
        )
        if result.anomalies:
            for index, anomaly in enumerate(result.anomalies, start=1):
                issue = anomaly.get("issue", "Issue requires clarification")
                action = anomaly.get("action", "Please clarify and submit evidence.")
                doc.add_paragraph(f"{index}. {issue} Required action: {action}")
        else:
            doc.add_paragraph("No material anomaly was automatically detected; manual credit review remains mandatory.")
        doc.add_paragraph("\nAssessment Officer                              Regional Relationship Manager")
        doc.save(path)
        return path


def make_result_zip(case_dir: Path, output_zip: Path) -> Path:
    with zipfile.ZipFile(output_zip, "w", zipfile.ZIP_DEFLATED) as archive:
        for path in sorted((case_dir / "outputs").rglob("*")):
            if path.is_file():
                archive.write(path, path.relative_to(case_dir / "outputs"))
    return output_zip
